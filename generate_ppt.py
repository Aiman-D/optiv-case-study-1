from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def add_text_slide(prs, title_text, content_text):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = title_text
    
    ORANGE = RGBColor(237, 107, 45)
    if slide.shapes.title.text_frame.paragraphs:
        slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = ORANGE

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "" 
    
    lines = content_text.split('\n')
    for idx, line in enumerate(lines):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
            
        line_stripped = line.lstrip()
        indent_level = 0
        if line.startswith("    "):
            indent_level = 2
        elif line.startswith("  "):
            indent_level = 1
            
        p.text = line_stripped
        p.level = indent_level
        p.font.size = Pt(20)

def add_box(slide, text, left, top, width, height, fill_color=RGBColor(20, 43, 76), font_size=14):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = RGBColor(255, 255, 255)
    
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    return shape

def add_down_arrow(slide, left, top, height=0.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(left), Inches(top), Inches(0.2), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(237, 107, 45) # Orange
    shape.line.color.rgb = RGBColor(237, 107, 45)

def create_presentation():
    prs = Presentation()

    NAVY_BLUE = RGBColor(20, 43, 76)
    ORANGE = RGBColor(237, 107, 45)

    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "Expected Outcomes: Case Study 01\nSolution Guide & Walkthrough"
    title.text_frame.paragraphs[0].font.color.rgb = NAVY_BLUE
    subtitle.text = "Vendor AI Usage Analysis Solution\nOptiv Consulting"

    # Introduction
    add_text_slide(prs, "Introduction to the Solution", 
        "Welcome to the Vendor AI Usage Analysis Tool Guide.\n"
        "This presentation will walk you through how our solution automatically assesses if third-party vendors are using AI.\n"
        "Key Principles:\n"
        "  • 100% OSINT (Public Information) - We don't contact the vendor.\n"
        "  • Verbatim Evidence - We only use exact quotes as proof.\n"
        "  • Tiered Approach - We spend more time on riskier vendors."
    )

    # Slide 1: Functional Design Diagram (Visual)
    slide_layout = prs.slide_layouts[5] # Blank with title
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "01 Functional Design Diagram"
    title.text_frame.paragraphs[0].font.color.rgb = ORANGE
    
    # Draw flowchart
    center_x = 4.0
    width = 3.5
    height = 0.6
    
    add_box(slide, "1. Input: Vendor Excel File", center_x - width/2, 1.5, width, height)
    add_down_arrow(slide, center_x - 0.1, 2.1, 0.4)
    
    add_box(slide, "2. M1: Criticality Engine (Scoring)", center_x - width/2, 2.5, width, height)
    add_down_arrow(slide, center_x - 0.1, 3.1, 0.4)
    
    add_box(slide, "3. M2: OSINT Collectors (Web Crawl)", center_x - width/2, 3.5, width, height)
    add_down_arrow(slide, center_x - 0.1, 4.1, 0.4)
    
    add_box(slide, "4. M3: AI Signal Classifier (LLM)", center_x - width/2, 4.5, width, height)
    add_down_arrow(slide, center_x - 0.1, 5.1, 0.4)
    
    add_box(slide, "5. M4 & M5: Risk Engine & Excel Output", center_x - width/2, 5.5, width, height)

    # Slide 2: Architecture Diagram (Visual)
    slide_layout = prs.slide_layouts[5] # Blank with title
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "02 Architecture Diagram"
    title.text_frame.paragraphs[0].font.color.rgb = ORANGE

    box_w = 2.0
    box_h = 1.0
    
    # Draw architecture blocks
    add_box(slide, "App Interface\n(Streamlit UI)", 4.0, 1.2, box_w, box_h, RGBColor(0,0,0))
    
    # Modules
    add_box(slide, "M1: config/rubric.yaml\n(pandas, yaml)", 1.0, 2.8, box_w*1.2, box_h)
    add_box(slide, "M2: trafilatura\n(urllib, local cache)", 4.0, 2.8, box_w*1.2, box_h)
    add_box(slide, "M3: LLM prompt\n(regex, JSON schema)", 7.0, 2.8, box_w*1.2, box_h)
    
    add_box(slide, "M4: Formula weighting\n(base × tier × recency)", 2.5, 4.5, box_w*1.5, box_h)
    add_box(slide, "M5: shutil.copy\n(pandas, openpyxl)", 6.0, 4.5, box_w*1.5, box_h)
    
    # Slide 3: Source Strategy
    add_text_slide(prs, "03 Source Strategy",
        "Where we look for evidence and how we prioritize it:\n"
        "Tier A (Direct Evidence):\n"
        "  The most reliable. Product documentation, AI policies, sub-processor lists.\n"
        "Tier B (Corroborating Evidence):\n"
        "  Technical job postings, engineering blogs, GitHub repositories.\n"
        "Tier C (Contextual Evidence):\n"
        "  Executive posts, press releases, interviews.\n"
        "Tier D (Noise):\n"
        "  Generic marketing buzzwords (we give this very little weight).\n"
        "Scoring System:\n"
        "  We multiply the reliability of the source by how closely it links to the actual service being provided."
    )

    # Slide 4: Criticality Ratings
    add_text_slide(prs, "04 Criticality Ratings",
        "How we decide which vendors need the most scrutiny:\n"
        "The Scoring Rubric:\n"
        "  We score vendors 0-3 on things like data sensitivity, payment flow, and operational dependency.\n"
        "The Resulting Tiers:\n"
        "  Critical: Gets a 'Full Sweep' (docs, job postings, financial filings, exec channels).\n"
        "  High: Checks product docs, trust center, jobs, and filings.\n"
        "  Medium: Checks website, basic docs, and jobs.\n"
        "  Low: Just a quick check of their trust center and AI policy.\n"
        "Why this matters: It saves time and resources by focusing on the highest risk vendors."
    )

    # Slide 5: Working Demonstration
    add_text_slide(prs, "05 Working Demonstration",
        "How to use the solution yourself:\n"
        "The User Interface:\n"
        "  We built a simple dashboard. You don't need to be a programmer to use it.\n"
        "The Process:\n"
        "  1. Upload your 'Meridian_Vendor_Input.xlsx' file into the app.\n"
        "  2. Click 'Run Pipeline'.\n"
        "  3. Watch as the tool shows 'evidence cards' for each vendor in real-time.\n"
        "  4. Download the final, completed Excel workbook with all the answers filled in.\n"
        "Evidence Log:\n"
        "  The final Excel file will have a new tab showing every exact quote and website link we found."
    )

    prs.save('Optiv_Engagement_Diagrams.pptx')
    print("Presentation created successfully as 'Optiv_Engagement_Diagrams.pptx'")

if __name__ == '__main__':
    create_presentation()
